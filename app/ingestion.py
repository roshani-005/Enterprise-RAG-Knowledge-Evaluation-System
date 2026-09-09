from __future__ import annotations

import io
import uuid
from dataclasses import dataclass
from pathlib import Path

import fitz
import pytesseract
from docx import Document as DocxDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from PIL import Image

from app.config import settings


@dataclass
class ParsedDocument:
    pages: list[Document]
    used_ocr: bool


def _extract_pdf(content: bytes, source: str) -> ParsedDocument:
    pdf = fitz.open(stream=content, filetype='pdf')
    pages: list[Document] = []
    used_ocr = False

    for page_number, page in enumerate(pdf, start=1):
        text = page.get_text('text').strip()
        if len(text) < settings.native_text_min_chars:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes('png')))
            ocr_text = pytesseract.image_to_string(image).strip()
            if len(ocr_text) > len(text):
                text = ocr_text
                used_ocr = True

        if text:
            pages.append(Document(page_content=text, metadata={'source': source, 'page': page_number}))

    return ParsedDocument(pages=pages, used_ocr=used_ocr)


def _extract_docx(content: bytes, source: str) -> ParsedDocument:
    doc = DocxDocument(io.BytesIO(content))
    text = '\n'.join(p.text for p in doc.paragraphs if p.text.strip())
    pages = [Document(page_content=text, metadata={'source': source, 'page': 1})] if text else []
    return ParsedDocument(pages=pages, used_ocr=False)


def _extract_txt(content: bytes, source: str) -> ParsedDocument:
    text = content.decode('utf-8', errors='replace').strip()
    pages = [Document(page_content=text, metadata={'source': source, 'page': 1})] if text else []
    return ParsedDocument(pages=pages, used_ocr=False)


def parse_document(filename: str, content: bytes) -> ParsedDocument:
    suffix = Path(filename).suffix.lower()
    if suffix == '.pdf':
        return _extract_pdf(content, filename)
    if suffix == '.docx':
        return _extract_docx(content, filename)
    if suffix in {'.txt', '.md'}:
        return _extract_txt(content, filename)
    raise ValueError(f'Unsupported file type: {suffix or "unknown"}')


def chunk_documents(pages: list[Document]) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=settings.chunk_size,
        chunk_overlap=settings.chunk_overlap,
        separators=['\n\n', '\n', '. ', ' ', ''],
    )
    chunks = splitter.split_documents(pages)
    for chunk in chunks:
        chunk.metadata['chunk_id'] = str(uuid.uuid4())
    return chunks
